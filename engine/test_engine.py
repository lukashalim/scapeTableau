from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest

from engine import transpile, transpile_formula
from engine.schema import ParseStatus, WorkbookParseError

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "samples" / "superstore_calcs.twb"


def test_iif_to_case_when() -> None:
    sql, status, err = transpile_formula('IIF([Sales] > 0, [Profit], 0)')
    assert status == ParseStatus.ok
    assert err is None
    assert sql == 'CASE WHEN "Sales" > 0 THEN "Profit" ELSE 0 END'


def test_zn_to_coalesce() -> None:
    sql, status, err = transpile_formula("ZN([Profit])")
    assert status == ParseStatus.ok
    assert err is None
    assert sql == 'COALESCE("Profit", 0)'


def test_ifnull_to_coalesce() -> None:
    sql, status, _ = transpile_formula("IFNULL([Discount], 0)")
    assert status == ParseStatus.ok
    assert sql == 'COALESCE("Discount", 0)'


def test_contains_literal_like() -> None:
    sql, status, _ = transpile_formula('CONTAINS([Product Name], "Table")')
    assert status == ParseStatus.ok
    assert sql == "\"Product Name\" LIKE '%Table%'"


def test_contains_expression_concat() -> None:
    sql, status, _ = transpile_formula("CONTAINS([Product Name], [Category])")
    assert status == ParseStatus.ok
    assert sql == "\"Product Name\" LIKE '%' || \"Category\" || '%'"


def test_lod_fixed_window() -> None:
    formula = "{FIXED [Region] : SUM([Sales])}"
    sql, status, err = transpile_formula(formula)
    assert status == ParseStatus.ok
    assert err is None
    assert sql == 'SUM("Sales") OVER (PARTITION BY "Region")'


def test_lod_fixed_no_dims() -> None:
    sql, status, _ = transpile_formula("{FIXED : SUM([Sales])}")
    assert status == ParseStatus.ok
    assert sql == 'SUM("Sales") OVER ()'


def test_include_without_view_grain() -> None:
    sql, status, err = transpile_formula("{INCLUDE [Category] : SUM([Sales])}")
    assert status == ParseStatus.skipped
    assert err == "INCLUDE/EXCLUDE requires view grain"
    assert "-- SKIPPED [lod-include/exclude]:" in sql


def test_iif_and_zn_dax() -> None:
    dax, status, err = transpile_formula("IIF([Sales] > 0, [Profit], 0)", target="dax")
    assert status == ParseStatus.ok
    assert err is None
    assert dax == "IF(Orders[Sales] > 0, Orders[Profit], 0)"
    zn, zn_status, _ = transpile_formula("ZN([Profit])", target="dax")
    assert zn_status == ParseStatus.ok
    assert zn == "COALESCE(Orders[Profit], 0)"


def test_lod_dax_calculate_allexcept() -> None:
    dax, status, err = transpile_formula("{FIXED [Region] : SUM([Sales])}", target="dax")
    assert status == ParseStatus.ok
    assert err is None
    assert dax == "CALCULATE(SUM(Orders[Sales]), ALLEXCEPT(Orders, Orders[Region]))"


def test_max_fixed_none_is_calculate_not_maxx() -> None:
    dax, status, err = transpile_formula("MAX({FIXED : SUM([Sales])})", target="dax")
    assert status == ParseStatus.ok
    assert err is None
    assert dax == "CALCULATE(SUM(Orders[Sales]), ALL(Orders))"
    assert "MAXX" not in dax


def test_dateadd_month_emits_edate() -> None:
    dax, status, err = transpile_formula("DATEADD('month', 1, [Order Date])", target="dax")
    assert status == ParseStatus.ok
    assert err is None
    assert dax == "EDATE(Orders[Order Date], 1)"


def test_sql_param_is_bare_literal() -> None:
    from engine.transpiler import FormulaContext
    from engine.schema import ParameterInfo

    ctx = FormulaContext(
        dialect="duckdb",
        parameters={"Multiplier": ParameterInfo(name="[Multiplier]", caption="Multiplier", datatype="real", value="2")},
    )
    sql, status, err = transpile_formula("[Sales] * [Parameters].[Multiplier]", dialect="duckdb", context=ctx)
    assert status == ParseStatus.ok, err
    assert sql == '"Sales" * 2'
    assert "/*" not in sql


def test_if_elseif_end() -> None:
    sql, status, _ = transpile_formula(
        'IF [Sales] > 1000 THEN "High" ELSEIF [Sales] > 500 THEN "Medium" ELSE "Low" END'
    )
    assert status == ParseStatus.ok
    assert sql == (
        'CASE WHEN "Sales" > 1000 THEN \'High\' '
        'WHEN "Sales" > 500 THEN \'Medium\' ELSE \'Low\' END'
    )


def test_unknown_function_does_not_crash() -> None:
    sql, status, err = transpile_formula("SOMEUNKNOWN([Sales])")
    assert status == ParseStatus.unparseable
    assert "SOMEUNKNOWN" in (err or "")
    assert "TODO: unmapped Tableau function" in sql


def test_sample_workbook_round_trip() -> None:
    result = transpile(SAMPLE)
    captions = {field.caption for field in result.calculated_fields}
    assert "Profit Ratio" in captions
    assert "Profit or Zero" in captions
    assert "Has Table" in captions
    assert "Regional Sales" in captions
    assert "West Profit" in captions

    zn = next(f for f in result.calculated_fields if f.caption == "Profit or Zero")
    assert zn.sql == 'COALESCE("Profit", 0)'
    assert zn.raw_formula == zn.formula
    assert zn.sql_expression == zn.sql
    assert zn.status == ParseStatus.ok
    dumped = json.loads(result.model_dump_json())
    first = dumped["calculated_fields"][0]
    assert "formula" in first and "sql" in first
    assert "dax" in first
    assert "raw_formula" not in first
    assert "sql_expression" not in first

    contains = next(f for f in result.calculated_fields if f.caption == "Has Table")
    assert contains.sql == "\"Product Name\" LIKE '%Table%'"

    lod = next(f for f in result.calculated_fields if f.caption == "Regional Sales")
    assert lod.status == ParseStatus.ok
    assert 'SUM("Sales") OVER (PARTITION BY "Region")' == lod.sql
    assert lod.dax == "CALCULATE(SUM(Orders[Sales]), ALLEXCEPT(Orders, Orders[Region]))"
    assert lod.dax_status == ParseStatus.ok

    resolved = next(f for f in result.calculated_fields if f.caption == "Ratio Abs")
    assert 'ABS((SUM("Profit") / SUM("Sales")))' == resolved.sql
    assert "Calculation_" not in resolved.sql
    assert result.dax_script
    assert "{{ config" not in result.sql_script

    ds_names = {ds.caption or ds.name for ds in result.datasources}
    assert "Sample Superstore" in ds_names
    assert "West Custom SQL" in ds_names
    assert any(ds.name.lower() == "parameters" for ds in result.datasources)

    dumped = json.loads(result.model_dump_json())
    blob = json.dumps(dumped)
    assert "s3cret" not in blob
    assert "hunter2" not in blob
    west = next(ds for ds in result.datasources if ds.caption == "West Custom SQL")
    assert any(conn.password == "***" for conn in west.connections)
    assert any(conn.server and "PWD=***" in conn.server for conn in west.connections)
    assert any(rel.type == "text" and rel.custom_sql and "West" in rel.custom_sql for rel in west.relations)

    hidden = [f for f in (result.inspect.fields if result.inspect else []) if f.hidden]
    assert any(f.caption == "Internal Order Id" for f in hidden)
    assert result.inspect and any(p.caption == "Profit Goal" for p in result.inspect.parameters)
    assert any(d.name == "Overview Dash" for d in result.dashboards)
    assert dumped.get("inspect")
    assert dumped.get("extract_inventory")

    assert "CREATE VIEW" in result.sql_script
    # One view per datasource (not one per calc); row calcs as columns.
    assert result.sql_script.count("CREATE VIEW") == 2
    assert 'CREATE VIEW "Sample Superstore" AS' in result.sql_script
    assert 'CREATE VIEW "West Custom SQL" AS' in result.sql_script
    assert 'CREATE VIEW "Profit Ratio" AS' not in result.sql_script
    assert 'COALESCE("Profit", 0) AS "Profit or Zero"' in result.sql_script
    assert 'SUM("Sales") OVER (PARTITION BY "Region") AS "Regional Sales"' in result.sql_script
    assert "-- Aggregates / LOD" in result.sql_script
    assert '-- "Profit Ratio":' in result.sql_script
    assert '-- "Ratio Abs":' in result.sql_script
    assert "-- SKIPPED [unparseable] Mystery Metric" in result.sql_script
    ratio = next(f for f in result.calculated_fields if f.caption == "Profit Ratio")
    assert ratio.sql_kind == "aggregate"
    zn_kind = next(f for f in result.calculated_fields if f.caption == "Profit or Zero")
    assert zn_kind.sql_kind == "row"
    lod_kind = next(f for f in result.calculated_fields if f.caption == "Regional Sales")
    assert lod_kind.sql_kind == "row"
    assert "Profit Ratio" in result.markdown_audit
    assert "| Caption |" in result.markdown_audit


def test_worksheet_table_calcs_listed_as_skipped() -> None:
    """Worksheet <table-calc> pills must appear in calculated_fields even without SQL windows."""

    result = transpile(SAMPLE)
    captions = {field.caption: field for field in result.calculated_fields}
    assert "% of Total Total Profit" in captions
    assert "% of Total Total Revenue" in captions

    profit = captions["% of Total Total Profit"]
    assert profit.formula == "% of Total SUM(Total Profit) along Table (Across)"
    assert profit.status == ParseStatus.skipped
    assert "table calc" in (profit.error or "").lower()
    assert "-- SKIPPED [table calc]" in profit.sql
    assert "Table (Across)" in profit.sql
    # Stage 2: DAX % of Total approximation while SQL remains skipped.
    assert profit.dax_status == ParseStatus.ok
    assert "DIVIDE(" in (profit.dax or "")
    assert "ALLSELECTED" in (profit.dax or "")

    revenue = captions["% of Total Total Revenue"]
    assert revenue.formula == "% of Total SUM(Total Revenue) along Table (Across)"
    assert revenue.status == ParseStatus.skipped
    assert revenue.dax_status == ParseStatus.ok

    dumped = json.loads(result.model_dump_json())
    dumped_captions = {f["caption"] for f in dumped["calculated_fields"]}
    assert "% of Total Total Profit" in dumped_captions
    assert "% of Total Total Revenue" in dumped_captions
    assert "-- SKIPPED [skipped] % of Total Total Profit" in result.sql_script

    # Inspect field list includes table-calc pills as calculated columns.
    assert result.inspect is not None
    inspect_captions = {f.caption for f in result.inspect.fields if f.calculated}
    assert "% of Total Total Profit" in inspect_captions
    assert "% of Total Total Revenue" in inspect_captions
    assert result.inspect.counts.calculated_fields >= len(result.calculated_fields)


def test_classify_sql_kind_row_vs_aggregate() -> None:
    from engine.transpiler import classify_sql_kind, parse_formula

    assert classify_sql_kind(parse_formula("ZN([Profit])")) == "row"
    assert classify_sql_kind(parse_formula("SUM([Sales])")) == "aggregate"
    assert classify_sql_kind(parse_formula("ABS(SUM([Profit])/SUM([Sales]))")) == "aggregate"
    assert (
        classify_sql_kind(parse_formula("{FIXED [Region]: SUM([Sales])}")) == "row"
    )


def test_dbt_script_one_view_per_datasource() -> None:
    result = transpile(SAMPLE, script_style="dbt")
    assert result.sql_script.count("{{ config(materialized='view') }}") == 2
    assert result.sql_script.count("CREATE VIEW") == 2


def test_twbx_bytes_round_trip() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("superstore_calcs.twb", SAMPLE.read_bytes())
    result = transpile(buf.getvalue())
    assert any(f.caption == "ZN" or f.caption == "Profit or Zero" for f in result.calculated_fields)


def test_invalid_xml_raises() -> None:
    with pytest.raises(WorkbookParseError, match="Invalid Tableau XML"):
        transpile(b"<not-tableau")


def test_invalid_zip_raises() -> None:
    with pytest.raises(WorkbookParseError, match="Invalid"):
        transpile(b"PK\x03\x04not-a-real-zip")


def test_asgi_app_post_transpile() -> None:
    import asyncio

    from app import app as asgi_app

    payload = SAMPLE.read_bytes()
    boundary = "----scapeBoundary"
    body = (
        f"--{boundary}\r\n"
        'Content-Disposition: form-data; name="file"; filename="superstore_calcs.twb"\r\n'
        "Content-Type: application/xml\r\n\r\n"
    ).encode("utf-8") + payload + f"\r\n--{boundary}--\r\n".encode("utf-8")

    messages = [{"type": "http.request", "body": body, "more_body": False}]

    async def receive() -> dict:
        return messages.pop(0)

    sent: list[dict] = []

    async def send(message: dict) -> None:
        sent.append(message)

    scope = {
        "type": "http",
        "method": "POST",
        "path": "/api/transpile",
        "headers": [(b"content-type", f"multipart/form-data; boundary={boundary}".encode())],
    }
    asyncio.run(asgi_app(scope, receive, send))
    assert sent[0]["status"] == 200
    data = json.loads(sent[1]["body"].decode("utf-8"))
    assert len(data["calculated_fields"]) >= 1
    assert "formula" in data["calculated_fields"][0]


def test_vercel_fn_json_keys_match_api() -> None:
    from engine.vercel_fn import transpile_payload

    status, payload = transpile_payload("superstore_calcs.twb", SAMPLE.read_bytes())
    assert status == 200
    body = json.loads(payload.decode("utf-8"))
    field = body["calculated_fields"][0]
    assert "formula" in field
    assert "sql" in field
    assert "raw_formula" not in field
    assert "sql_expression" not in field


def test_hidden_parameter_dashboard_on_sample() -> None:
    from engine.parser import parse_workbook

    parsed = parse_workbook(SAMPLE)
    assert any(f.hidden and f.caption == "Internal Order Id" for f in parsed.fields)
    assert any(p.caption == "Profit Goal" for p in parsed.parameters)
    assert parsed.dashboards and parsed.dashboards[0].zones


def test_csv_to_xlsx_and_hyper_listed_not_converted() -> None:
    from openpyxl import load_workbook

    from engine.extract import extract_to_xlsx

    twbx = ROOT / "samples" / "superstore_calcs.twbx"
    xlsx, inventory = extract_to_xlsx(twbx)
    assert any(m.kind == "csv" and m.path.endswith("Orders.csv") for m in inventory.archive_members)
    book = load_workbook(io.BytesIO(xlsx), read_only=True)
    names = book.sheetnames
    book.close()
    assert any("Orders" in name for name in names)
    assert any(name == "Custom SQL" for name in names)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("packaged.twb", SAMPLE.read_bytes())
        zf.writestr("extracts/sales.hyper", b"HYPERFAKE")
    xlsx2, inv2 = extract_to_xlsx(buf.getvalue())
    assert any(m.kind == "hyper" and m.path.endswith("sales.hyper") for m in inv2.hyper_files)
    book2 = load_workbook(io.BytesIO(xlsx2), read_only=True)
    try:
        assert "Hyper files" in book2.sheetnames
        rows = list(book2["Hyper files"].iter_rows(values_only=True))
        joined = " ".join(str(c) for row in rows for c in row)
        assert "sales.hyper" in joined
        assert "HYPERFAKE" not in joined
        assert "Listed only" in joined
    finally:
        book2.close()
    assert all("sales.hyper" != sheet for sheet in inv2.converted_sheets)


def test_twbx_hyper_sheet_uses_relation_name(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TWB relation name wins over Hyper Extract / TableauTemp table ids."""
    pytest.importorskip("tableauhyperapi")
    from openpyxl import load_workbook

    from engine.extract import extract_to_xlsx
    from engine.hyper import hyper_conversion_enabled

    monkeypatch.delenv("VERCEL", raising=False)
    if not hyper_conversion_enabled():
        pytest.skip("tableauhyperapi unavailable")

    hyper_path = ROOT / "samples" / "tiny_extract.hyper"
    twb = """<?xml version='1.0' encoding='utf-8' ?>
<workbook>
  <datasources>
    <datasource caption='Amazon Sales' inline='true' name='federated.sales'>
      <connection class='hyper' dbname='Data/Extracts/Extract.hyper' />
      <relation name='Orders' table='[Extract].[Extract]' type='table' />
      <column caption='Region' name='[Region]' role='dimension' />
    </datasource>
  </datasources>
</workbook>
"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("sales.twb", twb)
        zf.writestr("Data/Extracts/Extract.hyper", hyper_path.read_bytes())
    class _Named(io.BytesIO):
        name = "sales.twbx"

    xlsx, inv = extract_to_xlsx(_Named(buf.getvalue()))
    assert "Orders" in inv.converted_sheets
    assert all("TableauTemp" not in s and "Extract_Extract" not in s for s in inv.converted_sheets)
    book = load_workbook(io.BytesIO(xlsx), read_only=True)
    try:
        assert "Orders" in book.sheetnames
        rows = list(book["Orders"].iter_rows(values_only=True))
        assert rows[0][:2] == ("Region", "Sales")
    finally:
        book.close()


def test_convert_cli_sidecar_writes_xlsx(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pytest.importorskip("tableauhyperapi")
    from engine.hyper import hyper_conversion_enabled

    monkeypatch.delenv("VERCEL", raising=False)
    if not hyper_conversion_enabled():
        pytest.skip("tableauhyperapi unavailable")

    from desktop.convert_cli import main as convert_main

    hyper_path = ROOT / "samples" / "tiny_extract.hyper"
    out = tmp_path / "out.xlsx"
    code = convert_main(["convert", str(hyper_path), str(out)])
    assert code == 0
    assert out.is_file() and out.stat().st_size > 0


def test_convert_cli_transpile_sql_one_view_per_datasource(capsys: pytest.CaptureFixture[str]) -> None:
    from desktop.convert_cli import main as convert_main

    code = convert_main(["transpile", str(SAMPLE), "--dialect", "ansi"])
    assert code == 0
    out = capsys.readouterr().out.strip().splitlines()[-1]
    payload = json.loads(out)
    assert payload["ok"] is True
    script = payload["sql_script"]
    assert "CREATE VIEW" in script
    assert 'CREATE VIEW "Sample Superstore" AS' in script
    assert 'CREATE VIEW "Profit Ratio" AS' not in script
    assert "-- Aggregates / LOD" in script
    assert payload.get("skipped")


def test_hyper_sheet_title_skips_tableau_temp() -> None:
    from engine.hyper import (
        HyperTableRows,
        preferred_sheet_names_for_hyper,
        sheet_title_for,
    )
    from engine.schema import ConnectionInfo, DatasourceInfo, RelationInfo

    assert sheet_title_for(
        "Data/Extracts/Extract.hyper",
        "Extract",
        "#TableauTemp_0rpujo703as6kc17x6",
    ) == "Extract"
    assert (
        sheet_title_for(
            "Data/Extracts/sales_extract.hyper",
            "Extract",
            "TableauTemp_abc",
        )
        == "sales_extract"
    )
    assert (
        sheet_title_for(
            "x.hyper",
            "Extract",
            "#TableauTemp_x",
            preferred="Orders",
        )
        == "Orders"
    )

    ds = DatasourceInfo(
        name="federated.1",
        caption="Amazon Sales",
        connections=[
            ConnectionInfo(
                **{"class": "hyper", "dbname": "Data/Extracts/Extract.hyper"}
            )
        ],
        relations=[
            RelationInfo(
                name="Orders",
                type="table",
                table="[Extract].[#TableauTemp_0rpujo703as6kc17x6]",
            )
        ],
    )
    tables = [
        HyperTableRows(
            schema="Extract",
            table="#TableauTemp_0rpujo703as6kc17x6",
            columns=["a"],
            rows=[],
        )
    ]
    prefs = preferred_sheet_names_for_hyper(
        [ds], "Data/Extracts/Extract.hyper", tables
    )
    assert prefs == ["Orders"]
    assert (
        sheet_title_for(
            "Data/Extracts/Extract.hyper",
            tables[0].schema,
            tables[0].table,
            preferred=prefs[0],
        )
        == "Orders"
    )


def test_hyper_date_timestamp_cell_values_excel_safe() -> None:
    """tableauhyperapi Date has no isoformat — must become datetime.date for openpyxl."""
    pytest.importorskip("tableauhyperapi")
    from datetime import date, datetime

    from openpyxl import Workbook
    from tableauhyperapi import Date, Timestamp

    from engine.hyper import _cell_value

    d = _cell_value(Date(2010, 5, 28))
    t = _cell_value(Timestamp(2010, 5, 28, 12, 30, 0))
    assert d == date(2010, 5, 28)
    assert isinstance(t, datetime)
    assert t.replace(tzinfo=None) == datetime(2010, 5, 28, 12, 30, 0)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append([d, t])  # must not raise IllegalWorksheetData


def test_standalone_hyper_and_tdsx_extract(monkeypatch: pytest.MonkeyPatch) -> None:
    from openpyxl import load_workbook

    from engine.extract import extract_to_xlsx
    from engine.hyper import hyper_conversion_enabled
    from engine.schema import WorkbookParseError

    monkeypatch.delenv("VERCEL", raising=False)
    hyper_path = ROOT / "samples" / "tiny_extract.hyper"
    assert hyper_path.is_file()

    if hyper_conversion_enabled():
        xlsx, inv = extract_to_xlsx(hyper_path)
        assert any(m.kind == "hyper" for m in inv.hyper_files)
        # Standalone .hyper → zip member stem (not Extract_Extract)
        assert "tiny_extract" in inv.converted_sheets
        assert all("TableauTemp" not in s for s in inv.converted_sheets)
        book = load_workbook(io.BytesIO(xlsx), read_only=True)
        try:
            data_sheets = [n for n in book.sheetnames if n != "Hyper files"]
            assert data_sheets == ["tiny_extract"]
            rows = list(book[data_sheets[0]].iter_rows(values_only=True))
            assert rows[0][:2] == ("Region", "Sales")
        finally:
            book.close()

    # .tdsx = zip with hyper, no .twb — extract must not require TWB
    tdsx = io.BytesIO()
    with zipfile.ZipFile(tdsx, "w") as zf:
        zf.writestr("Data/Extracts/tiny.hyper", hyper_path.read_bytes())
    tdsx_bytes = tdsx.getvalue()
    # filename matters for kind detection when calling via path-like isn't available
    class _NamedBytes(io.BytesIO):
        name = "pack.tdsx"

    named = _NamedBytes(tdsx_bytes)
    if hyper_conversion_enabled():
        xlsx2, inv2 = extract_to_xlsx(named)
        assert inv2.hyper_files
        assert "tiny" in inv2.converted_sheets
        assert all("TableauTemp" not in s for s in inv2.converted_sheets)

    # Transpile still rejects TDSX without TWB
    from engine import transpile

    with pytest.raises(WorkbookParseError):
        transpile(named)


def test_tiny_hyper_converts_locally_skipped_on_vercel(monkeypatch: pytest.MonkeyPatch) -> None:
    from openpyxl import load_workbook

    from engine.extract import extract_to_xlsx
    from engine.hyper import hyper_conversion_enabled

    hyper_path = ROOT / "samples" / "tiny_extract.hyper"
    assert hyper_path.is_file()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("packaged.twb", SAMPLE.read_bytes())
        zf.writestr("Data/Extracts/tiny_extract.hyper", hyper_path.read_bytes())
    payload = buf.getvalue()

    monkeypatch.delenv("VERCEL", raising=False)
    if hyper_conversion_enabled():
        xlsx, inv = extract_to_xlsx(payload)
        # TWB has CSV relations only — do not steal "Orders"; use hyper stem.
        assert "tiny_extract" in inv.converted_sheets
        assert "Orders" not in inv.converted_sheets
        book = load_workbook(io.BytesIO(xlsx), read_only=True)
        try:
            rows = list(book["tiny_extract"].iter_rows(values_only=True))
            assert rows[0][:2] == ("Region", "Sales")
            regions = {row[0] for row in rows[1:]}
            assert regions == {"West", "East"}
        finally:
            book.close()
        assert any("converted locally" in n.lower() for n in inv.notes)

    monkeypatch.setenv("VERCEL", "1")
    xlsx_v, inv_v = extract_to_xlsx(payload)
    assert "tiny_extract" not in inv_v.converted_sheets
    assert any(
        "Browser extract cannot convert Hyper" in n or "desktop" in n.lower() or "listed only" in n.lower()
        for n in inv_v.notes
    )
    book_v = load_workbook(io.BytesIO(xlsx_v), read_only=True)
    try:
        assert "Hyper files" in book_v.sheetnames
        assert "tiny_extract" not in book_v.sheetnames
    finally:
        book_v.close()


def test_password_redaction_connect_string() -> None:
    from engine.redact import redact_connect_fragment, redact_password_attr

    assert redact_password_attr("s3cret") == "***"
    assert redact_connect_fragment("host;PWD=hunter2;Database=x") == "host;PWD=***;Database=x"
    assert redact_connect_fragment("Password=abc;foo=1") == "Password=***;foo=1"


def test_asgi_app_post_extract() -> None:
    import asyncio

    from app import app as asgi_app

    payload = (ROOT / "samples" / "superstore_calcs.twbx").read_bytes()
    boundary = "----scapeBoundary"
    body = (
        f"--{boundary}\r\n"
        'Content-Disposition: form-data; name="file"; filename="superstore_calcs.twbx"\r\n'
        "Content-Type: application/zip\r\n\r\n"
    ).encode("utf-8") + payload + f"\r\n--{boundary}--\r\n".encode("utf-8")

    messages = [{"type": "http.request", "body": body, "more_body": False}]

    async def receive() -> dict:
        return messages.pop(0)

    sent: list[dict] = []

    async def send(message: dict) -> None:
        sent.append(message)

    scope = {
        "type": "http",
        "method": "POST",
        "path": "/api/extract",
        "headers": [(b"content-type", f"multipart/form-data; boundary={boundary}".encode())],
    }
    asyncio.run(asgi_app(scope, receive, send))
    assert sent[0]["status"] == 200
    headers = {k.decode(): v.decode() for k, v in sent[0]["headers"]}
    assert "spreadsheetml.sheet" in headers["content-type"]
    assert sent[1]["body"][:2] == b"PK"


def test_cli_inspect_and_extract(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from cli import app

    runner = CliRunner()
    inspected = runner.invoke(app, ["inspect", str(SAMPLE), "--format", "json"])
    assert inspected.exit_code == 0, inspected.output
    payload = json.loads(inspected.stdout)
    assert payload["counts"]["parameters"] >= 1
    assert "s3cret" not in inspected.stdout

    out = tmp_path / "extract.xlsx"
    extracted = runner.invoke(app, ["extract", str(ROOT / "samples" / "superstore_calcs.twbx"), "--output", str(out)])
    assert extracted.exit_code == 0, extracted.output
    assert out.exists() and out.read_bytes()[:2] == b"PK"


def test_cli_json_stdin(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from typer.testing import CliRunner

    from cli import app

    runner = CliRunner()
    payload = SAMPLE.read_bytes()
    result = runner.invoke(app, ["process", "-", "--format", "json"], input=payload)
    assert result.exit_code == 0, result.output
    assert '"calculated_fields"' in result.stdout

    out_dir = tmp_path / "out"
    written = runner.invoke(
        app,
        ["process", str(SAMPLE), "--output-dir", str(out_dir), "--format", "all", "--print-summary"],
    )
    assert written.exit_code == 0, written.output
    assert (out_dir / "escapetableau.sql").exists()
    assert (out_dir / "audit.md").exists()
    assert (out_dir / "result.json").exists()
    assert "Profit Ratio" in written.output
