"""In-memory .twbx/.tdsx/.hyper extract → openpyxl workbook bytes."""

from __future__ import annotations

import csv
import io
import os
import zipfile
from pathlib import Path
from typing import BinaryIO
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

from engine.hyper import (
    MAX_ROWS,
    hyper_conversion_enabled,
    read_hyper_bytes,
    sheet_title_for,
)
from engine.parser import (
    Source,
    _is_zip,
    _local_tag,
    _read_bytes,
    is_hyper_bytes,
    list_archive_members,
    parse_workbook,
)
from engine.schema import ArchiveMember, CustomSqlEntry, ExtractInventory, WorkbookParseError

HYPER_NOTE_VERCEL = (
    "Hyper extracts are listed only on Vercel / the hosted site. "
    "Download the free BIEscape desktop converter for Excel on your computer."
)
HYPER_NOTE_LOCAL_OK = (
    "Hyper extracts were converted locally via tableauhyperapi "
    "(temp file for hyperd, deleted after read; 100k row cap)."
)
HYPER_NOTE_LOCAL_SKIP = (
    "Hyper extracts are listed only. Install optional tableauhyperapi to convert locally, "
    "or use the BIEscape desktop converter (not used when VERCEL is set)."
)


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = "".join("_" if ch in r'[]:*?/\\' else ch for ch in name).strip() or "Sheet"
    cleaned = cleaned[:31]
    candidate = cleaned
    n = 1
    while candidate in used:
        suffix = f"_{n}"
        candidate = (cleaned[: 31 - len(suffix)] + suffix)[:31]
        n += 1
    used.add(candidate)
    return candidate


def _write_rows(ws, rows: list[list[object]]) -> None:
    for row in rows[:MAX_ROWS]:
        ws.append(list(row))


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _csv_rows(data: bytes, *, delimiter: str) -> list[list[object]]:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: list[list[object]] = []
    for i, row in enumerate(reader):
        if i >= MAX_ROWS:
            break
        rows.append(list(row))
    return rows


def _xlsx_rows(data: bytes) -> dict[str, list[list[object]]]:
    bio = io.BytesIO(data)
    try:
        wb = load_workbook(bio, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — old .xls and corrupt books
        raise WorkbookParseError(f"Unable to read spreadsheet: {exc}") from exc
    sheets: dict[str, list[list[object]]] = {}
    try:
        for ws in wb.worksheets:
            rows: list[list[object]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    break
                rows.append(["" if cell is None else cell for cell in row])
            sheets[ws.title] = rows
    finally:
        wb.close()
    return sheets


def build_inventory_from_parts(
    members: list[ArchiveMember],
    custom_sql: list[CustomSqlEntry],
    *,
    converted_sheets: list[str] | None = None,
    extra_notes: list[str] | None = None,
) -> ExtractInventory:
    notes = list(extra_notes or [])
    hyper = [m for m in members if m.kind == "hyper"]
    tabular = [m for m in members if m.kind in {"csv", "xlsx"}]
    if not members:
        notes.append("This workbook has no packaged archive members (plain .twb).")
    elif not tabular and not custom_sql and hyper and not any(
        s for s in (converted_sheets or []) if "Extract" in s or "hyper" in s.lower()
    ):
        notes.append(
            "This package is Hyper-only for data files; use the BIEscape desktop converter "
            "or convert locally with tableauhyperapi."
        )
    elif not tabular and not custom_sql and not hyper:
        notes.append("No embedded CSV/TSV/XLS/XLSX tables were found to convert.")
    return ExtractInventory(
        archive_members=members,
        tabular_files=tabular,
        hyper_files=hyper,
        custom_sql=custom_sql,
        converted_sheets=list(converted_sheets or []),
        notes=notes,
    )


def build_inventory(source: Source | None = None, *, workbook=None) -> ExtractInventory:
    parsed = workbook if workbook is not None else parse_workbook(source)  # type: ignore[arg-type]
    custom = [
        CustomSqlEntry(name=rel.name, datasource=ds.caption or ds.name, sql=rel.custom_sql)
        for ds in parsed.datasources
        for rel in ds.relations
        if rel.custom_sql
    ]
    return build_inventory_from_parts(list(parsed.archive_members), custom)


def _custom_sql_from_twb(twb_bytes: bytes) -> list[CustomSqlEntry]:
    try:
        root = ET.fromstring(twb_bytes)
    except ET.ParseError:
        return []
    entries: list[CustomSqlEntry] = []
    for ds_el in root.iter():
        if _local_tag(ds_el.tag) != "datasource":
            continue
        ds_caption = ds_el.get("caption") or ds_el.get("name") or ""
        for rel in ds_el.iter():
            if _local_tag(rel.tag) != "relation":
                continue
            if (rel.get("type") or "").lower() != "text":
                continue
            sql = (rel.text or "").strip()
            if not sql:
                continue
            entries.append(
                CustomSqlEntry(name=rel.get("name"), datasource=ds_caption, sql=sql)
            )
    return entries


def _new_workbook():
    wb = Workbook()
    used: set[str] = set()
    default = wb.active
    assert default is not None
    converted: list[str] = []
    first_sheet_used = False

    def next_ws(title: str):
        nonlocal first_sheet_used
        name = _safe_sheet_name(title, used)
        if not first_sheet_used:
            default.title = name
            first_sheet_used = True
            return default
        return wb.create_sheet(name)

    return wb, default, used, converted, next_ws, lambda: first_sheet_used


def _append_hyper_payload(
    *,
    path: str,
    payload: bytes,
    size: int,
    next_ws,
    converted: list[str],
    notes: list[str],
    index_rows: list[list[object]],
) -> bool:
    """Convert one Hyper blob; return True if any table sheet was written."""

    can_convert = hyper_conversion_enabled()
    if not can_convert:
        index_rows.append(
            [
                path,
                size,
                "Listed only — not converted (Vercel or missing tableauhyperapi)",
            ]
        )
        return False
    result = read_hyper_bytes(payload)
    if result.error and not result.tables:
        index_rows.append([path, size, f"Listed only — {result.error}"])
        notes.append(f"{path}: {result.error}")
        return False
    wrote = False
    for table in result.tables:
        title = sheet_title_for(path, table.schema, table.table)
        ws = next_ws(title)
        if table.columns:
            ws.append(list(table.columns))
        _write_rows(ws, table.rows)
        converted.append(ws.title)
        wrote = True
        note = f"Converted ({table.schema}.{table.table})"
        if table.truncated:
            note += f"; truncated at {MAX_ROWS} rows"
            notes.append(f"{path}: truncated at {MAX_ROWS} rows")
        index_rows.append([path, size, note])
    if result.error:
        notes.append(f"{path}: {result.error}")
    return wrote


def extract_to_xlsx(source: str | Path | bytes | BinaryIO) -> tuple[bytes, ExtractInventory]:
    data, filename = _read_bytes(source)
    if not data:
        raise WorkbookParseError("Workbook is empty")

    # Standalone .hyper — skip XML / zip parse entirely
    if is_hyper_bytes(data, filename) and not _is_zip(data, filename):
        path = filename or "extract.hyper"
        member = ArchiveMember(path=path, size=len(data), kind="hyper")
        wb, _default, _used, converted, next_ws, first_used = _new_workbook()
        notes: list[str] = []
        index_rows: list[list[object]] = [["Path", "Size", "Note"]]
        hyper_converted = _append_hyper_payload(
            path=path,
            payload=data,
            size=len(data),
            next_ws=next_ws,
            converted=converted,
            notes=notes,
            index_rows=index_rows,
        )
        ws = next_ws("Hyper files")
        for row in index_rows:
            ws.append(row)
        converted.append(ws.title)
        if hyper_converted:
            notes.append(HYPER_NOTE_LOCAL_OK)
        else:
            notes.append(HYPER_NOTE_VERCEL if os.environ.get("VERCEL") else HYPER_NOTE_LOCAL_SKIP)
        if not first_used():
            pass
        inventory = build_inventory_from_parts(
            [member], [], converted_sheets=converted, extra_notes=notes
        )
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), inventory

    zip_bytes = data if _is_zip(data, filename) else None
    members: list[ArchiveMember] = []
    custom: list[CustomSqlEntry] = []

    if zip_bytes is not None:
        members = list_archive_members(zip_bytes)
        twb_members = [m for m in members if m.kind == "twb"]
        if twb_members:
            try:
                with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
                    twb_bytes = zf.read(twb_members[0].path)
                # Prefer full parse when TWB exists (custom SQL + captions)
                try:
                    parsed = parse_workbook(source)
                    members = list(parsed.archive_members) or members
                    custom = [
                        CustomSqlEntry(
                            name=rel.name,
                            datasource=ds.caption or ds.name,
                            sql=rel.custom_sql,
                        )
                        for ds in parsed.datasources
                        for rel in ds.relations
                        if rel.custom_sql
                    ]
                except WorkbookParseError:
                    custom = _custom_sql_from_twb(twb_bytes)
            except KeyError:
                pass
        # .tdsx / zip without TWB: members from archive listing only
    else:
        # Plain .twb
        parsed = parse_workbook(data)
        members = list(parsed.archive_members)
        custom = [
            CustomSqlEntry(name=rel.name, datasource=ds.caption or ds.name, sql=rel.custom_sql)
            for ds in parsed.datasources
            for rel in ds.relations
            if rel.custom_sql
        ]

    wb, default, used, converted, next_ws, first_used = _new_workbook()
    notes = []

    if zip_bytes is not None:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            for member in members:
                if member.kind not in {"csv", "xlsx"}:
                    continue
                payload = zf.read(member.path)
                stem = Path(member.path).name
                lower = member.path.lower()
                try:
                    if lower.endswith(".csv"):
                        rows = _csv_rows(payload, delimiter=",")
                        ws = next_ws(stem)
                        _write_rows(ws, rows)
                        converted.append(ws.title)
                    elif lower.endswith(".tsv"):
                        rows = _csv_rows(payload, delimiter="\t")
                        ws = next_ws(stem)
                        _write_rows(ws, rows)
                        converted.append(ws.title)
                    else:
                        sheets = _xlsx_rows(payload)
                        for sheet_name, rows in sheets.items():
                            ws = next_ws(f"{Path(stem).stem}_{sheet_name}")
                            _write_rows(ws, rows)
                            converted.append(ws.title)
                except WorkbookParseError as exc:
                    notes.append(f"{member.path}: {exc.message}")

    if custom:
        ws = next_ws("Custom SQL")
        ws.append(["Datasource", "Relation", "SQL"])
        for entry in custom:
            ws.append([entry.datasource or "", entry.name or "", entry.sql])
        converted.append(ws.title)

    hyper = [m for m in members if m.kind == "hyper"]
    hyper_converted = False
    if hyper:
        index_rows: list[list[object]] = [["Path", "Size", "Note"]]
        if zip_bytes is not None:
            with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
                for item in hyper:
                    payload = zf.read(item.path)
                    if _append_hyper_payload(
                        path=item.path,
                        payload=payload,
                        size=item.size,
                        next_ws=next_ws,
                        converted=converted,
                        notes=notes,
                        index_rows=index_rows,
                    ):
                        hyper_converted = True
        else:
            for item in hyper:
                index_rows.append(
                    [
                        item.path,
                        item.size,
                        "Listed only — Hyper member without archive bytes",
                    ]
                )
        ws = next_ws("Hyper files")
        for row in index_rows:
            ws.append(row)
        converted.append(ws.title)
        if hyper_converted:
            notes.append(HYPER_NOTE_LOCAL_OK)
        else:
            notes.append(HYPER_NOTE_VERCEL if os.environ.get("VERCEL") else HYPER_NOTE_LOCAL_SKIP)

    if not first_used():
        default.title = _safe_sheet_name("Extract", used)
        default.append(["Note"])
        default.append(["No embedded CSV/TSV/XLS/XLSX tables were converted."])
        converted.append(default.title)

    inventory = build_inventory_from_parts(members, custom, converted_sheets=converted, extra_notes=notes)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), inventory
